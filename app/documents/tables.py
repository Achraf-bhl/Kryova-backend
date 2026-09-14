"""Spreadsheets, kept as tables: CSV, TSV and Excel workbooks (master plan P4.2).

*"Spreadsheets keep their structure -- a load-case table becomes rows with
units, not prose."* So every row is a `TABLE_ROW` fragment whose `cells` each
carry their own `sheet`/`cell` locator and the heading of the column they sit
under. "Cell C7 of loads.xlsx" is then sayable about one value three steps
downstream (P4.4), and a unit written in a heading -- `Fx [N]` -- travels with
every value under it.

**How far a value is believed follows the file, not the look of the value.**
A workbook types its cells, so a numeric, boolean or date cell is `MEASURED`: the
file itself asserts that number. A string is `TRANSCRIBED`, even when it spells
a number -- somebody typed `1200` as text, and the reader will say it parses
(`Cell.number`) without promoting it. A CSV has no types at all, so every cell
in one is `TRANSCRIBED`.

Four ways a spreadsheet misleads a reader that walks it naively, and what this
one does about each:

* **A formula with no saved result has no value.** A workbook written by a
  program that does not calculate -- openpyxl, pandas, most exporters -- stores
  `=B2*2` and nothing else, and read for values the cell is simply *absent*:
  measured on openpyxl 3.1.5, the cell does not appear in the row at all. That
  is a load case with a force missing and nothing to say so. It is recorded as
  `Unread`, quoting the formula, with how to get a value.
* **A formula with a saved result is `TRANSCRIBED`, not `MEASURED`.** The number
  is the result of a calculation this reader did not repeat, saved by whatever
  last wrote the file, and nothing here can tell whether it is current.
* **An error is not a string.** `#DIV/0!` in a load table is recorded as
  `Unread`, not quoted as a value a consumer might skip.
* **Hidden sheets, rows and columns are read, and said.** They are the
  workbook's version of a PDF's `/Title`: content the person who sent the file
  may never have seen. Dropping them would hide an injection channel instead of
  labelling it; reading them silently would present a superseded load case,
  hidden rather than deleted, as current.

**The workbook is loaded whole, under a size budget, never in read-only mode.**
openpyxl's read-only worksheets have no row or column dimensions and no merged
ranges (measured: `AttributeError` on `row_dimensions`), so hiddenness is
unknowable there -- and hiddenness is the thing worth knowing. The declared
unpacked size is checked first (`structure.MAX_PACKAGE_BYTES`). And the cells
are walked from the ones that exist, never from the sheet's bounding box: one
cell at XFD1048576 makes that box seventeen billion cells.
"""

from __future__ import annotations

import csv
import io
import logging
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from app.documents.document import DocumentKind, ExtractedDocument, Fragment, FragmentKind, Unread
from app.documents.errors import ExtractionFailed
from app.documents.provenance import Reliability, SourceRef
from app.documents.quoted import UntrustedText
from app.documents.structure import (
    MAX_PACKAGE_BYTES,
    Collector,
    RawCell,
    add_table,
    column_letter,
    decode_text,
    open_package,
    package_properties,
    strict_number,
    unread_package_contents,
)

logger = logging.getLogger(__name__)

#: The delimiters a CSV is tried against. Comma and tab are what the suffix
#: promises; semicolon is what Excel writes wherever the decimal separator is a
#: comma; the pipe turns up in exports from older systems.
_DELIMITERS: tuple[str, ...] = (",", ";", "\t", "|")

#: How many records decide the delimiter. Enough to see a consistent column
#: count, few enough that a malformed tail cannot change the decision.
_DELIMITER_SAMPLE_RECORDS = 20


# -- CSV and TSV ----------------------------------------------------------------


def read_delimited(
    path: Path, source: SourceRef, label: str, *, max_fragments: int, max_text_bytes: int
) -> ExtractedDocument:
    """A delimited text table, one `TABLE_ROW` per record.

    Rows are numbered by *record*, not by line, so a quoted cell with a newline
    in it does not shift every row after it -- the number is the one a
    spreadsheet shows when it opens the file.
    """
    size = path.stat().st_size
    if size > max_text_bytes:
        raise ExtractionFailed(
            f"{source.filename} is {size / 1e6:.0f} MB, over the "
            f"{max_text_bytes / 1e6:.0f} MB limit for a table read as text. Attach the "
            "rows that matter instead.",
            short="too large to read as a table",
        )

    text, encoding_note = decode_text(path.read_bytes())
    document_source = source.read_by("csv", Reliability.TRANSCRIBED)
    collector = Collector(limit=max_fragments, what="the table")
    if encoding_note is not None:
        collector.notes.append(encoding_note)

    default = "\t" if label == "tsv" else ","
    delimiter, delimiter_note = _choose_delimiter(text, default)
    if delimiter_note is not None:
        collector.notes.append(delimiter_note)

    rows: list[tuple[SourceRef, list[RawCell]]] = []
    try:
        for number, record in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), 1):
            if len(rows) > max_fragments:
                break
            cells: list[RawCell] = []
            for column, value in enumerate(record, start=1):
                if not value.strip():
                    continue
                parsed = strict_number(value)
                cells.append(
                    RawCell(
                        column=column,
                        text=value,
                        source=document_source.at(cell=f"{column_letter(column)}{number}"),
                        number=parsed,
                        is_text=parsed is None,
                    )
                )
            rows.append((document_source.at(row=number), cells))
    except csv.Error as exc:
        raise ExtractionFailed(
            f"{source.filename} could not be read as a table ({exc}). Open it in a "
            "spreadsheet, save it as CSV UTF-8, and attach that.",
            short="not a readable table",
        ) from exc

    heading = add_table(rows, collector)
    if heading is not None:
        collector.notes.append(
            f"Row {heading.locator.row} was taken as the column headings."
        )

    return ExtractedDocument(
        source=document_source,
        kind=DocumentKind.SPREADSHEET,
        fragments=tuple(collector.fragments),
        notes=tuple(collector.notes),
    )


def _choose_delimiter(text: str, default: str) -> tuple[str, str | None]:
    """The delimiter, and a note when the file did not settle it.

    **A rule, not a sniffer.** A candidate is *consistent* when it splits every
    sampled record into the same number of columns, and more than one. Exactly
    one consistent candidate is used. When several are -- `LC1;1,5;2,3` is three
    columns split on either -- the file genuinely does not say, so the suffix's
    delimiter is used and the ambiguity is named, because a decimal comma read
    as a column break turns one force into two plausible numbers.
    `csv.Sniffer` breaks that same tie by a fixed preference for the comma and
    says nothing, which is the wrong half of the outcome.
    """
    consistent: list[str] = []
    for candidate in _DELIMITERS:
        widths: set[int] = set()
        try:
            for index, record in enumerate(csv.reader(io.StringIO(text), delimiter=candidate)):
                if index >= _DELIMITER_SAMPLE_RECORDS:
                    break
                if any(value.strip() for value in record):
                    widths.add(len(record))
        except csv.Error:
            continue
        if len(widths) == 1 and next(iter(widths)) > 1:
            consistent.append(candidate)

    names = {",": "commas", ";": "semicolons", "\t": "tabs", "|": "pipes"}
    if len(consistent) == 1:
        chosen = consistent[0]
        if chosen == default:
            return chosen, None
        return chosen, f"The columns are separated by {names[chosen]}, and were read that way."
    if len(consistent) > 1:
        alternatives = " or ".join(names[c] for c in consistent)
        chosen = default if default in consistent else consistent[0]
        return chosen, (
            f"The columns could be split on {alternatives}; they were read as separated "
            f"by {names[chosen]}. If the columns look wrong -- a decimal comma read as "
            "a column break -- re-save the file with a different separator."
        )
    return default, None


# -- Excel workbooks ------------------------------------------------------------


def read_workbook(path: Path, source: SourceRef, *, max_fragments: int) -> ExtractedDocument:
    """Every sheet of an `.xlsx`/`.xlsm` as rows of cells, plus its properties.

    Loaded twice -- once for the saved values, once for the formulas -- because
    openpyxl gives one or the other per load, and telling "a formula with no
    saved result" from "an empty cell" needs both.
    """
    try:
        import openpyxl  # noqa: PLC0415 - imported where used, like ezdxf and pypdf
    except ImportError as exc:  # pragma: no cover - openpyxl is in requirements.txt
        raise ExtractionFailed(
            f"{source.filename} is an Excel workbook, but openpyxl is not installed in "
            "this deployment. Install openpyxl, or save the sheet as CSV and attach that.",
            short="openpyxl not installed",
        ) from exc

    document_source = source.read_by("openpyxl", Reliability.TRANSCRIBED)
    collector = Collector(limit=max_fragments, what="the workbook")

    archive = open_package(path, source, max_bytes=MAX_PACKAGE_BYTES)
    try:
        package_properties(archive, document_source, collector)
        unread_package_contents(archive, collector, root="xl")
    finally:
        archive.close()

    try:
        # A file handle, not the path: openpyxl refuses a path whose extension
        # is not .xlsx/.xlsm, and the blob store names files by digest.
        with path.open("rb") as handle:
            saved = openpyxl.load_workbook(handle, data_only=True)
        with path.open("rb") as handle:
            written = openpyxl.load_workbook(handle, data_only=False)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises a wide family
        raise ExtractionFailed(
            f"{source.filename}: openpyxl could not read this workbook "
            f"({type(exc).__name__}: {exc}). If it is a binary workbook (.xlsb), save it "
            "as .xlsx; otherwise open it and save it again, then attach that.",
            short="the workbook could not be read",
        ) from exc

    sheet_count = len(written.worksheets)
    for index, sheet in enumerate(written.worksheets, start=1):
        if collector.full:
            break
        _read_sheet(
            sheet,
            saved[sheet.title],
            document_source.at(sheet=sheet.title),
            collector,
            label=f"Sheet {index}" if sheet_count > 1 else "The sheet",
        )

    if written.chartsheets:
        collector.notes.append(
            f"{len(written.chartsheets)} chart sheet(s) were not read -- a chart is a "
            "picture of numbers, and the numbers are on the sheet it plots."
        )

    return ExtractedDocument(
        source=document_source,
        kind=DocumentKind.SPREADSHEET,
        fragments=tuple(collector.fragments),
        unread=tuple(collector.unread),
        notes=tuple(collector.notes),
    )


def _read_sheet(
    written: Any, saved: Any, sheet_source: SourceRef, collector: Collector, *, label: str
) -> None:
    """One worksheet. `label` is "Sheet 2", never the sheet's name.

    Notes are server-authored prose rendered outside the fence, and a sheet name
    is text whoever made the workbook chose -- so the notes count sheets and the
    name travels only inside the locators, where it is quoted.
    """
    state = getattr(written, "sheet_state", "visible")
    if state == "hidden":
        collector.notes.append(
            f"{label} is hidden in the workbook and was read anyway -- the person who "
            "sent it may never have seen it."
        )
    elif state == "veryHidden":
        collector.notes.append(
            f"{label} is 'very hidden' -- it can only be shown by a macro -- and was read "
            "anyway. Treat what it says with suspicion."
        )

    # `_cells` is the dict of cells the sheet actually holds. `iter_rows` walks
    # the bounding box instead, creating every empty cell in it.
    held = sorted(written._cells.items())  # noqa: SLF001 - see the module docstring
    if not held:
        collector.notes.append(f"{label} is empty.")
        return

    hidden_rows = sorted(row for row, dim in written.row_dimensions.items() if dim.hidden)
    hidden_columns = sorted(
        {
            column
            for dim in written.column_dimensions.values()
            if dim.hidden and dim.min and dim.max
            for column in range(dim.min, dim.max + 1)
        }
    )
    if hidden_rows or hidden_columns:
        what = []
        if hidden_rows:
            what.append(
                ("row " if len(hidden_rows) == 1 else "rows ")
                + ", ".join(str(row) for row in hidden_rows[:12])
            )
        if hidden_columns:
            what.append(
                ("column " if len(hidden_columns) == 1 else "columns ")
                + ", ".join(column_letter(column) for column in hidden_columns[:12])
            )
        collector.notes.append(
            f"{label} has hidden {' and '.join(what)}. Hidden cells were read anyway -- "
            "the person who sent it may not have seen them, and hidden rows in a load "
            "table are often superseded cases."
        )
    if written.merged_cells.ranges:
        collector.notes.append(
            f"{label} has merged cells; a merged value is read once, at the top-left "
            "cell of its range."
        )

    by_row: dict[int, list[RawCell]] = {}
    for (row, column), cell in held:
        coordinate = f"{column_letter(column)}{row}"
        where = sheet_source.at(cell=coordinate)

        comment = getattr(cell, "comment", None)
        if comment is not None and str(comment.text).strip():
            collector.add(
                Fragment(
                    text=UntrustedText(f"Comment: {comment.text}", where),
                    kind=FragmentKind.ANNOTATION,
                )
            )

        raw = _cell(cell, saved.cell(row=row, column=column), where, collector)
        if raw is not None:
            by_row.setdefault(row, []).append(raw)

    rows = [(sheet_source.at(row=row), cells) for row, cells in sorted(by_row.items())]
    heading = add_table(rows, collector)
    if heading is not None:
        collector.notes.append(
            f"{label}: row {heading.locator.row} was taken as the column headings."
        )


def _cell(written: Any, saved: Any, where: SourceRef, collector: Collector) -> RawCell | None:
    """One cell as a `RawCell`, or None -- having recorded why when it matters."""
    column = int(written.column)
    is_formula = written.data_type == "f"
    value = saved.value if is_formula else written.value

    if is_formula and value is None:
        collector.unread.append(
            Unread(
                where=where,
                what=_formula_text(written.value),
                why=(
                    "this cell holds a formula with no saved result -- the workbook was "
                    "last written by a program that does not calculate, so there is no "
                    "value to read. Open it in Excel or LibreOffice, save it, and attach "
                    "it again"
                ),
            )
        )
        return None
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if (saved.data_type if is_formula else written.data_type) == "e":
        collector.unread.append(
            Unread(
                where=where,
                what=str(value),
                why="the cell shows a spreadsheet error instead of a value",
            )
        )
        return None

    typed = Reliability.TRANSCRIBED if is_formula else Reliability.MEASURED
    if isinstance(value, bool):
        return RawCell(
            column, "TRUE" if value else "FALSE", where.read_by("openpyxl", typed), None, False
        )
    if isinstance(value, (int, float)):
        return RawCell(
            column, _number_text(value), where.read_by("openpyxl", typed), float(value), False
        )
    if isinstance(value, datetime):
        text = value.date().isoformat() if value.time() == time(0) else value.isoformat()
        return RawCell(column, text, where.read_by("openpyxl", typed), None, False)
    if isinstance(value, (date, time)):
        return RawCell(column, value.isoformat(), where.read_by("openpyxl", typed), None, False)
    if isinstance(value, timedelta):
        return RawCell(column, str(value), where.read_by("openpyxl", typed), None, False)

    text = str(value)
    return RawCell(
        column,
        text,
        where.read_by("openpyxl", Reliability.TRANSCRIBED),
        strict_number(text),
        True,
    )


def _number_text(value: int | float) -> str:
    """The number as stored: an integer without `.0`, a float at full precision.

    `repr` rather than a display format, because the cell's *format* -- two
    decimals, a percentage -- is how one screen showed it, and the stored value
    is what the file asserts.
    """
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return repr(value)


def _formula_text(formula: Any) -> str:
    """A formula as written, including an array formula's `text`."""
    text = getattr(formula, "text", formula)
    return str(text)


__all__ = ["read_delimited", "read_workbook"]
